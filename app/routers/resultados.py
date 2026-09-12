from __future__ import annotations

from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.models.campeonato import Campeonato
from app.models.categoria import Categoria
from app.models.fecha import Fecha
from app.models.jinete import Jinete
from app.models.jinete_campeonato import JineteCampeonato
from app.models.resultado import ResultadoCategoria, ResultadoDetalle
from app.models.sorteo import Sorteo, SorteoDetalle
from app.services.resultados import posiciones_campeonato


templates = Jinja2Templates(directory="app/templates")

router = APIRouter(prefix="/sorteos/resultados", tags=["Resultados"])
public_router = APIRouter(prefix="/publico/resultados", tags=["Resultados públicos"])

ROLES_LECTURA = {"MASTER", "ADMIN", "SECRETARIA", "LOCUCION", "TV"}
ROLES_EDICION = {"MASTER", "ADMIN"}


def exigir_lectura(request: Request) -> None:
    if str(request.session.get("usuario_rol") or "").upper() not in ROLES_LECTURA:
        raise HTTPException(status_code=403, detail="No tiene permisos para ver Resultados.")


def exigir_edicion(request: Request) -> None:
    if str(request.session.get("usuario_rol") or "").upper() not in ROLES_EDICION:
        raise HTTPException(status_code=403, detail="No tiene permisos para modificar Resultados.")


def detalle_utiles(sorteo: Sorteo | None) -> list[SorteoDetalle]:
    if sorteo is None:
        return []
    return [d for d in sorteo.detalles if not d.es_reserva and d.jinete_id is not None]


def obtener_resultado(db: Session, sorteo_id: int) -> ResultadoCategoria | None:
    return db.scalar(
        select(ResultadoCategoria).where(ResultadoCategoria.sorteo_id == sorteo_id)
    )


def contexto_seleccion(
    db: Session,
    campeonato_id: int | None,
    fecha_id: int | None,
    categoria_id: int | None,
):
    campeonatos = db.scalars(
        select(Campeonato)
        .join(Fecha, Fecha.campeonato_id == Campeonato.id)
        .join(Sorteo, Sorteo.fecha_id == Fecha.id)
        .distinct()
        .order_by(Campeonato.fecha_inicio.desc(), Campeonato.id.desc())
    ).all()
    if campeonato_id is None and campeonatos:
        campeonato_id = campeonatos[0].id

    fechas = []
    if campeonato_id is not None:
        fechas = db.scalars(
            select(Fecha)
            .join(Sorteo, Sorteo.fecha_id == Fecha.id)
            .where(Fecha.campeonato_id == campeonato_id)
            .distinct()
            .order_by(Fecha.fecha.asc(), Fecha.id.asc())
        ).all()
    if fecha_id not in {f.id for f in fechas}:
        fecha_id = fechas[0].id if fechas else None

    categorias = []
    if fecha_id is not None:
        categorias = db.scalars(
            select(Categoria)
            .join(Sorteo, Sorteo.categoria_id == Categoria.id)
            .where(Sorteo.fecha_id == fecha_id)
            .distinct()
            .order_by(Categoria.orden.asc(), Categoria.nombre.asc())
        ).all()
    if categoria_id not in {c.id for c in categorias}:
        categoria_id = categorias[0].id if categorias else None

    sorteo = None
    if fecha_id is not None and categoria_id is not None:
        sorteo = db.scalar(
            select(Sorteo).where(
                Sorteo.fecha_id == fecha_id,
                Sorteo.categoria_id == categoria_id,
            )
        )

    return campeonatos, fechas, categorias, sorteo, campeonato_id, fecha_id, categoria_id


def url_panel(sorteo: Sorteo | None) -> str:
    if sorteo is None:
        return "/sorteos/resultados"
    return (
        f"/sorteos/resultados?campeonato_id={sorteo.fecha.campeonato_id}"
        f"&fecha_id={sorteo.fecha_id}&categoria_id={sorteo.categoria_id}"
    )


def parse_puntos(valor) -> Decimal | None:
    if valor is None or str(valor).strip() == "":
        return None
    texto = str(valor).strip().replace(",", ".")
    try:
        numero = Decimal(texto)
    except InvalidOperation as exc:
        raise ValueError(f"Puntaje inválido: {valor}") from exc
    return numero.quantize(Decimal("0.01"))


def texto_celda(valor) -> str:
    return "" if valor is None else str(valor).strip()


@router.get("", response_class=HTMLResponse)
def panel_resultados(
    request: Request,
    campeonato_id: int | None = Query(default=None),
    fecha_id: int | None = Query(default=None),
    categoria_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
):
    exigir_lectura(request)
    campeonatos, fechas, categorias, sorteo, campeonato_id, fecha_id, categoria_id = contexto_seleccion(
        db, campeonato_id, fecha_id, categoria_id
    )
    resultado = obtener_resultado(db, sorteo.id) if sorteo else None

    por_detalle = {}
    if resultado:
        por_detalle = {d.sorteo_detalle_id: d for d in resultado.detalles}

    filas = []
    for detalle in detalle_utiles(sorteo):
        cargado = por_detalle.get(detalle.id)
        filas.append({
            "detalle": detalle,
            "puntos": cargado.puntos if cargado else None,
            "observaciones": cargado.observaciones if cargado else None,
        })

    posiciones = []
    if sorteo is not None and categoria_id is not None:
        posiciones = posiciones_campeonato(
            db,
            sorteo.fecha.campeonato_id,
            categoria_id,
            solo_publicados=False,
        )

    return templates.TemplateResponse(
        request=request,
        name="resultados/panel.html",
        context={
            "menu_activo": "resultados",
            "usuario_nombre": request.session.get("usuario_nombre", "Usuario"),
            "rol": str(request.session.get("usuario_rol") or "").upper(),
            "campeonatos": campeonatos,
            "fechas": fechas,
            "categorias": categorias,
            "campeonato_id": campeonato_id,
            "fecha_id": fecha_id,
            "categoria_id": categoria_id,
            "sorteo": sorteo,
            "resultado": resultado,
            "filas": filas,
            "posiciones": posiciones,
        },
    )


@router.get("/excel")
def exportar_excel(
    request: Request,
    sorteo_id: int = Query(...),
    db: Session = Depends(get_db),
):
    exigir_lectura(request)
    sorteo = db.get(Sorteo, sorteo_id)
    if sorteo is None:
        raise HTTPException(status_code=404, detail="Sorteo oficial no encontrado.")

    detalles = detalle_utiles(sorteo)
    if not detalles:
        raise HTTPException(status_code=400, detail="El sorteo no tiene montas oficiales para exportar.")

    resultado = obtener_resultado(db, sorteo.id)
    cargados = {d.sorteo_detalle_id: d for d in resultado.detalles} if resultado else {}

    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados"
    visibles = ["Orden", "Palenque", "Jinete", "Localidad", "Caballo", "Tropilla", "Puntos", "Observaciones"]
    tecnicas = ["_sorteo_id", "_sorteo_detalle_id", "_fecha_id", "_categoria_id", "_jinete_id", "_caballo_id"]
    ws.append(visibles + tecnicas)

    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="101820")
        celda.alignment = Alignment(horizontal="center", vertical="center")

    for detalle in detalles:
        existente = cargados.get(detalle.id)
        ws.append([
            detalle.orden,
            detalle.palenque,
            detalle.jinete_nombre or "",
            detalle.jinete_localidad or "",
            detalle.caballo_nombre or "",
            detalle.tropilla_nombre or "",
            float(existente.puntos) if existente and existente.puntos is not None else None,
            existente.observaciones if existente else "",
            sorteo.id,
            detalle.id,
            sorteo.fecha_id,
            sorteo.categoria_id,
            detalle.jinete_id,
            detalle.caballo_id,
        ])

    anchos = {"A": 10, "B": 12, "C": 30, "D": 24, "E": 30, "F": 26, "G": 14, "H": 42}
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho
    for col in "IJKLMN":
        ws.column_dimensions[col].hidden = True
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{ws.max_row}"

    archivo = BytesIO()
    wb.save(archivo)
    archivo.seek(0)
    nombre = f"resultados_fecha_{sorteo.fecha_id}_categoria_{sorteo.categoria_id}.xlsx"
    return StreamingResponse(
        archivo,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )


@router.post("/importar")
async def importar_excel(
    request: Request,
    sorteo_id: int = Form(...),
    archivo: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    exigir_edicion(request)
    sorteo = db.get(Sorteo, sorteo_id)
    if sorteo is None:
        raise HTTPException(status_code=404, detail="Sorteo oficial no encontrado.")

    resultado = obtener_resultado(db, sorteo.id)
    if resultado is not None and resultado.estado != "borrador":
        raise HTTPException(status_code=400, detail="El resultado está finalizado. Reabrilo antes de importar otra planilla.")

    contenido = await archivo.read()
    try:
        wb = load_workbook(BytesIO(contenido), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="No se pudo abrir el archivo Excel.") from exc
    ws = wb["Resultados"] if "Resultados" in wb.sheetnames else wb.active

    encabezados = [texto_celda(ws.cell(1, c).value) for c in range(1, 15)]
    esperados = ["Orden", "Palenque", "Jinete", "Localidad", "Caballo", "Tropilla", "Puntos", "Observaciones", "_sorteo_id", "_sorteo_detalle_id", "_fecha_id", "_categoria_id", "_jinete_id", "_caballo_id"]
    if encabezados != esperados:
        raise HTTPException(status_code=400, detail="La planilla no tiene el formato oficial de Resultados.")

    oficiales = {d.id: d for d in detalle_utiles(sorteo)}
    filas_excel: dict[int, tuple[Decimal | None, str | None]] = {}

    for fila in range(2, ws.max_row + 1):
        if all(ws.cell(fila, c).value is None for c in range(1, 15)):
            continue
        try:
            sid = int(ws.cell(fila, 9).value)
            detalle_id = int(ws.cell(fila, 10).value)
            fecha_id = int(ws.cell(fila, 11).value)
            categoria_id = int(ws.cell(fila, 12).value)
        except (TypeError, ValueError) as exc:
            raise HTTPException(status_code=400, detail=f"Fila {fila}: identificadores técnicos inválidos.") from exc

        if sid != sorteo.id or fecha_id != sorteo.fecha_id or categoria_id != sorteo.categoria_id:
            raise HTTPException(status_code=400, detail=f"Fila {fila}: la planilla corresponde a otro sorteo/fecha/categoría.")
        detalle = oficiales.get(detalle_id)
        if detalle is None:
            raise HTTPException(status_code=400, detail=f"Fila {fila}: la monta no pertenece al sorteo oficial.")
        if detalle_id in filas_excel:
            raise HTTPException(status_code=400, detail=f"Fila {fila}: la monta está duplicada.")

        jinete_id = ws.cell(fila, 13).value
        caballo_id = ws.cell(fila, 14).value
        if (int(jinete_id) if jinete_id is not None else None) != detalle.jinete_id or int(caballo_id) != detalle.caballo_id:
            raise HTTPException(status_code=400, detail=f"Fila {fila}: los identificadores de jinete/caballo fueron modificados.")

        visibles = [
            (1, str(detalle.orden)),
            (2, "" if detalle.palenque is None else str(detalle.palenque)),
            (3, detalle.jinete_nombre or ""),
            (4, detalle.jinete_localidad or ""),
            (5, detalle.caballo_nombre or ""),
            (6, detalle.tropilla_nombre or ""),
        ]
        for col, oficial in visibles:
            if texto_celda(ws.cell(fila, col).value) != texto_celda(oficial):
                raise HTTPException(status_code=400, detail=f"Fila {fila}: se modificó información oficial del sorteo ({esperados[col-1]}).")

        try:
            puntos = parse_puntos(ws.cell(fila, 7).value)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=f"Fila {fila}: {exc}") from exc
        obs = texto_celda(ws.cell(fila, 8).value) or None
        filas_excel[detalle_id] = (puntos, obs)

    if set(filas_excel) != set(oficiales):
        faltan = len(set(oficiales) - set(filas_excel))
        raise HTTPException(status_code=400, detail=f"La planilla no contiene todas las montas oficiales. Faltan {faltan} fila(s).")

    if resultado is None:
        resultado = ResultadoCategoria(sorteo_id=sorteo.id, estado="borrador")
        db.add(resultado)
        db.flush()
    else:
        for viejo in list(resultado.detalles):
            db.delete(viejo)
        db.flush()

    for detalle_id, (puntos, obs) in filas_excel.items():
        db.add(ResultadoDetalle(
            resultado_categoria_id=resultado.id,
            sorteo_detalle_id=detalle_id,
            puntos=puntos,
            observaciones=obs,
        ))
    resultado.actualizado_en = datetime.utcnow()
    db.commit()
    request.session["flash_success"] = "Planilla importada. Los resultados quedaron en borrador y no fueron publicados."
    return RedirectResponse(url_panel(sorteo), status_code=303)


@router.post("/{resultado_id}/finalizar")
def finalizar_resultado(request: Request, resultado_id: int, db: Session = Depends(get_db)):
    exigir_edicion(request)
    resultado = db.get(ResultadoCategoria, resultado_id)
    if resultado is None:
        raise HTTPException(status_code=404, detail="Resultado no encontrado.")
    if resultado.estado != "borrador":
        raise HTTPException(status_code=400, detail="Sólo un borrador puede finalizarse.")
    esperadas = len(detalle_utiles(resultado.sorteo))
    if len(resultado.detalles) != esperadas:
        raise HTTPException(status_code=400, detail="La carga no contiene todas las montas oficiales.")
    resultado.estado = "finalizado"
    resultado.finalizado_en = datetime.utcnow()
    resultado.actualizado_en = datetime.utcnow()
    db.flush()

    descalificados = aplicar_descalificacion_por_puntaje(
        db,
        resultado.sorteo.fecha.campeonato_id,
        resultado.sorteo.categoria_id,
    )
    db.commit()

    mensaje = "Categoría finalizada. Ya puede publicarse cuando corresponda."
    if descalificados:
        mensaje += (
            f" {descalificados} jinete(s) quedaron descalificados por no alcanzar "
            "5 puntos acumulados en las primeras 2 fechas."
        )
    request.session["flash_success"] = mensaje
    return RedirectResponse(url_panel(resultado.sorteo), status_code=303)


@router.post("/{resultado_id}/reabrir")
def reabrir_resultado(request: Request, resultado_id: int, db: Session = Depends(get_db)):
    exigir_edicion(request)
    resultado = db.get(ResultadoCategoria, resultado_id)
    if resultado is None:
        raise HTTPException(status_code=404, detail="Resultado no encontrado.")
    if resultado.estado != "finalizado":
        raise HTTPException(status_code=400, detail="Sólo un resultado finalizado y no publicado puede reabrirse.")
    resultado.estado = "borrador"
    resultado.finalizado_en = None
    resultado.actualizado_en = datetime.utcnow()
    db.commit()
    request.session["flash_success"] = "Resultado reabierto como borrador."
    return RedirectResponse(url_panel(resultado.sorteo), status_code=303)


@router.post("/{resultado_id}/publicar")
def publicar_resultado(request: Request, resultado_id: int, db: Session = Depends(get_db)):
    exigir_edicion(request)
    resultado = db.get(ResultadoCategoria, resultado_id)
    if resultado is None:
        raise HTTPException(status_code=404, detail="Resultado no encontrado.")
    if resultado.estado != "finalizado":
        raise HTTPException(status_code=400, detail="Primero tenés que finalizar la categoría.")
    resultado.estado = "publicado"
    resultado.publicado_en = datetime.utcnow()
    resultado.actualizado_en = datetime.utcnow()
    db.commit()
    request.session["flash_success"] = "Resultados publicados en el portal público."
    return RedirectResponse(url_panel(resultado.sorteo), status_code=303)


@router.post("/{resultado_id}/despublicar")
def despublicar_resultado(request: Request, resultado_id: int, db: Session = Depends(get_db)):
    exigir_edicion(request)
    resultado = db.get(ResultadoCategoria, resultado_id)
    if resultado is None:
        raise HTTPException(status_code=404, detail="Resultado no encontrado.")
    if resultado.estado != "publicado":
        raise HTTPException(status_code=400, detail="El resultado no está publicado.")
    resultado.estado = "finalizado"
    resultado.publicado_en = None
    resultado.actualizado_en = datetime.utcnow()
    db.commit()
    request.session["flash_success"] = "Resultados retirados del portal público."
    return RedirectResponse(url_panel(resultado.sorteo), status_code=303)



def aplicar_descalificacion_por_puntaje(
    db: Session,
    campeonato_id: int,
    categoria_id: int,
) -> int:
    """Descalifica si no alcanza 5 puntos acumulados en las primeras dos
    fechas oficiales de su categoría. Sólo corre cuando ambas están cerradas
    en Resultados (finalizado/publicado).
    """
    primeros_sorteos = db.scalars(
        select(Sorteo)
        .join(Fecha, Fecha.id == Sorteo.fecha_id)
        .where(
            Fecha.campeonato_id == campeonato_id,
            Sorteo.categoria_id == categoria_id,
        )
        .order_by(Fecha.fecha.asc(), Fecha.id.asc(), Sorteo.id.asc())
        .limit(2)
    ).all()

    if len(primeros_sorteos) < 2:
        return 0

    cargas = []
    for sorteo in primeros_sorteos:
        carga = obtener_resultado(db, sorteo.id)
        if carga is None or carga.estado not in {"finalizado", "publicado"}:
            return 0
        cargas.append(carga)

    puntos_por_jinete: dict[int, Decimal] = {}
    for carga in cargas:
        for rd in carga.detalles:
            detalle = rd.sorteo_detalle
            if detalle is None or detalle.es_reserva or detalle.jinete_id is None:
                continue
            puntos_por_jinete.setdefault(detalle.jinete_id, Decimal("0"))
            if rd.puntos is not None:
                puntos_por_jinete[detalle.jinete_id] += Decimal(rd.puntos)

    inscriptos = db.scalars(
        select(JineteCampeonato).where(
            JineteCampeonato.campeonato_id == campeonato_id,
            JineteCampeonato.categoria_id == categoria_id,
        )
    ).all()

    descalificados = 0
    for inscripto in inscriptos:
        total = puntos_por_jinete.get(inscripto.jinete_id, Decimal("0"))
        if total >= Decimal("5"):
            continue
        jinete = db.get(Jinete, inscripto.jinete_id)
        if jinete is not None and jinete.estado != "descalificado":
            jinete.estado = "descalificado"
            descalificados += 1

    return descalificados


def resultados_fecha_ordenados(resultado: ResultadoCategoria) -> list[dict]:
    filas = []
    for rd in resultado.detalles:
        d = rd.sorteo_detalle
        if d is None or d.es_reserva or d.jinete_id is None:
            continue
        puntos_decimal = Decimal(rd.puntos) if rd.puntos is not None else None
        filas.append({
            "orden": d.orden,
            "jinete": d.jinete_nombre or "-",
            "localidad": d.jinete_localidad or "-",
            "caballo": d.caballo_nombre or "-",
            "tropilla": d.tropilla_nombre or "-",
            "_puntos_decimal": puntos_decimal,
            "puntos": float(rd.puntos) if rd.puntos is not None else None,
            "observaciones": rd.observaciones or "",
        })

    filas.sort(
        key=lambda x: (
            x["_puntos_decimal"] is None,
            -(x["_puntos_decimal"] or Decimal("0")),
            str(x["jinete"]).casefold(),
        )
    )

    posicion = 0
    for fila in filas:
        if fila["_puntos_decimal"] is None:
            fila["posicion"] = None
        else:
            posicion += 1
            fila["posicion"] = posicion
        fila.pop("_puntos_decimal", None)
    return filas


@public_router.get("/datos", response_class=JSONResponse)
def datos_publicos(db: Session = Depends(get_db)):
    publicados = db.scalars(
        select(ResultadoCategoria)
        .join(Sorteo, Sorteo.id == ResultadoCategoria.sorteo_id)
        .join(Fecha, Fecha.id == Sorteo.fecha_id)
        .where(ResultadoCategoria.estado == "publicado")
        .order_by(Fecha.fecha.desc(), Fecha.id.desc())
    ).all()

    agrupado: dict[int, dict] = {}
    for resultado in publicados:
        sorteo = resultado.sorteo
        fecha = sorteo.fecha
        categoria = sorteo.categoria
        campeonato = fecha.campeonato
        camp = agrupado.setdefault(campeonato.id, {
            "id": campeonato.id,
            "nombre": campeonato.nombre,
            "categorias": {},
        })
        cat = camp["categorias"].setdefault(categoria.id, {
            "id": categoria.id,
            "nombre": categoria.nombre,
            "fechas": [],
            "posiciones": posiciones_campeonato(db, campeonato.id, categoria.id, solo_publicados=True),
        })
        cat["fechas"].append({
            "id": fecha.id,
            "nombre": fecha.nombre,
            "fecha": fecha.fecha.isoformat(),
            "localidad": fecha.localidad,
            "resultados": resultados_fecha_ordenados(resultado),
        })

    campeonatos = []
    for camp in agrupado.values():
        camp["categorias"] = list(camp["categorias"].values())
        campeonatos.append(camp)
    return JSONResponse({"campeonatos": campeonatos}, headers={"Cache-Control": "no-store"})
