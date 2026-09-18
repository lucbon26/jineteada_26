"""Carga independiente del repechaje previo a F3."""
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
from fastapi import APIRouter, Depends, Request, HTTPException, UploadFile, File
from fastapi.responses import RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import select
from openpyxl import Workbook, load_workbook
from app.core.database import get_db
from app.models.categoria import Categoria
from app.models.jinete import Jinete
from app.models.repechaje import RepechajeCategoria
from app.services.clasificacion import base_categoria, candidatos_repechaje, hay_sorteo_posterior, recalcular_categoria

router = APIRouter(prefix='/repechaje')
templates = Jinja2Templates(directory='app/templates')


def contexto(request, categoria_id, db, editar=False):
    rol = str(request.session.get('usuario_rol') or '').upper()
    if rol not in ({'MASTER', 'ADMIN'} if editar else {'MASTER', 'ADMIN', 'SECRETARIA', 'LOCUCION', 'TV'}):
        raise HTTPException(403, 'Sin permiso para esta operación.')
    cat = db.scalar(select(Categoria).where(Categoria.id == categoria_id).with_for_update())
    if not cat:
        raise HTTPException(404, 'Categoría no encontrada.')
    totales, firma = base_categoria(db, cat.campeonato_id, cat.id)
    jinetes = db.scalars(select(Jinete).where(Jinete.id.in_(candidatos_repechaje(totales))).order_by(Jinete.apellidos, Jinete.nombres)).all()
    planilla = db.get(RepechajeCategoria, cat.id)
    if editar and hay_sorteo_posterior(db, cat):
        raise HTTPException(409, 'Ya hay un sorteo de F3 o posterior. No se puede cambiar su clasificación; revisá primero ese sorteo.')
    return cat, totales, firma, jinetes, planilla


def puntaje(valor):
    if valor is None or str(valor).strip() == '':
        return None
    try:
        p = Decimal(str(valor).strip().replace(',', '.'))
        if not p.is_finite() or p < 0 or p > Decimal('99999999.99') or p != p.quantize(Decimal('.01')):
            raise ValueError()
        return str(p.quantize(Decimal('.01')))
    except (InvalidOperation, ValueError):
        raise HTTPException(400, 'Usá puntos no negativos, con hasta dos decimales. Vacío significa pendiente; 0 significa eliminado.')


def guardar(db, cat, totales, firma, jinetes, planilla, enviada, valores, confirmar, usuario):
    if totales is None or not jinetes:
        raise HTTPException(409, 'F1 y F2 deben estar completas y finalizadas, con participantes en repechaje.')
    if enviada != firma:
        raise HTTPException(409, 'Cambió la clasificación de F1/F2. Recargá la pantalla o exportá una planilla nueva.')
    if planilla and planilla.estado == 'confirmado' and planilla.base_firma == firma:
        raise HTTPException(409, 'Reabrí el repechaje antes de editarlo.')
    if set(valores) != {str(j.id) for j in jinetes}:
        raise HTTPException(400, 'La planilla debe contener exactamente los jinetes del repechaje, sin duplicados.')
    puntos = {jid: puntaje(v) for jid, v in valores.items()}
    if confirmar and any(v is None for v in puntos.values()):
        raise HTTPException(400, 'Completá todos los puntajes antes de confirmar; cargá 0 para quien no obtuvo puntos.')
    if planilla is None:
        planilla = RepechajeCategoria(categoria_id=cat.id)
        db.add(planilla)
    planilla.base_firma = firma
    planilla.puntos_json = json.dumps(puntos)
    planilla.estado = 'confirmado' if confirmar else 'borrador'
    planilla.confirmado_en = datetime.utcnow() if confirmar else None
    planilla.confirmado_por = usuario if confirmar else None
    db.flush()
    recalcular_categoria(db, cat.campeonato_id, cat.id)
    db.commit()


@router.get('/{categoria_id}')
def panel(request: Request, categoria_id: int, db=Depends(get_db)):
    cat, totales, firma, jinetes, planilla = contexto(request, categoria_id, db)
    vigente = planilla is not None and planilla.base_firma == firma and totales is not None
    return templates.TemplateResponse(request=request, name='resultados/repechaje.html', context={
        'categoria': cat, 'totales': totales, 'firma': firma, 'jinetes': jinetes,
        'planilla': planilla, 'vigente': vigente,
        'puntos': json.loads(planilla.puntos_json) if vigente else {},
        'bloqueado': hay_sorteo_posterior(db, cat),
        'edita': request.session.get('usuario_rol', '').upper() in {'MASTER', 'ADMIN'},
        'menu_activo': 'resultados',
    })


@router.post('/{categoria_id}')
async def cargar(request: Request, categoria_id: int, db=Depends(get_db)):
    cat, totales, firma, jinetes, planilla = contexto(request, categoria_id, db, True)
    form = await request.form()
    valores = {k[7:]: v for k, v in form.items() if k.startswith('puntos_')}
    guardar(db, cat, totales, firma, jinetes, planilla, form.get('firma'), valores, form.get('accion') == 'confirmar', request.session.get('usuario_nombre', 'Administrador'))
    request.session['flash_success'] = 'Repechaje confirmado; se actualizó quién puede seguir.' if form.get('accion') == 'confirmar' else 'Borrador guardado. El sorteo sigue bloqueado hasta confirmar.'
    return RedirectResponse(f'/sorteos/resultados/repechaje/{categoria_id}', 303)


@router.post('/{categoria_id}/reabrir')
def reabrir(request: Request, categoria_id: int, db=Depends(get_db)):
    cat, _, _, _, planilla = contexto(request, categoria_id, db, True)
    if planilla:
        planilla.estado = 'borrador'
        planilla.confirmado_en = None
        planilla.confirmado_por = None
        db.flush()
        recalcular_categoria(db, cat.campeonato_id, cat.id)
        db.commit()
    return RedirectResponse(f'/sorteos/resultados/repechaje/{categoria_id}', 303)


@router.get('/{categoria_id}/excel')
def exportar(request: Request, categoria_id: int, db=Depends(get_db)):
    cat, totales, firma, jinetes, planilla = contexto(request, categoria_id, db)
    if totales is None or not jinetes:
        raise HTTPException(409, 'No hay un repechaje disponible para exportar.')
    puntos = json.loads(planilla.puntos_json) if planilla and planilla.base_firma == firma else {}
    wb = Workbook(); ws = wb.active; ws.title = 'Repechaje'
    ws.append(['Jinete', 'Puntos repechaje', '_jinete_id', '_categoria_id', '_base_firma'])
    for j in jinetes:
        ws.append([f'{j.apellidos}, {j.nombres}', puntos.get(str(j.id)), j.id, cat.id, firma])
        ws.cell(ws.max_row, 1).data_type = 's'
    ws.column_dimensions['A'].width = 40; ws.column_dimensions['B'].width = 22
    for col in ['C', 'D', 'E']: ws.column_dimensions[col].hidden = True
    ws.freeze_panes = 'A2'
    out = BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(out, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename="repechaje_{cat.id}.xlsx"'})


@router.post('/{categoria_id}/importar')
async def importar(request: Request, categoria_id: int, archivo: UploadFile = File(...), db=Depends(get_db)):
    cat, totales, firma, jinetes, planilla = contexto(request, categoria_id, db, True)
    try:
        wb = load_workbook(BytesIO(await archivo.read()), data_only=True)
        ws = wb['Repechaje']
        if [c.value for c in ws[1]] != ['Jinete', 'Puntos repechaje', '_jinete_id', '_categoria_id', '_base_firma']:
            raise ValueError()
        nombres = {str(j.id): f'{j.apellidos}, {j.nombres}' for j in jinetes}
        valores = {}
        for nombre, puntos, jid, cid, base in ws.iter_rows(min_row=2, max_col=5, values_only=True):
            if all(v is None for v in [nombre, puntos, jid, cid, base]): continue
            clave = str(jid)
            if clave in valores or clave not in nombres or nombre != nombres[clave] or cid != cat.id or base != firma:
                raise ValueError()
            valores[clave] = puntos
    except Exception as exc:
        raise HTTPException(400, 'Planilla inválida o desactualizada. Exportá la plantilla de esta categoría y modificá sólo Puntos repechaje.') from exc
    guardar(db, cat, totales, firma, jinetes, planilla, firma, valores, False, '')
    request.session['flash_success'] = 'Excel importado como borrador. Revisalo y confirmá el repechaje.'
    return RedirectResponse(f'/sorteos/resultados/repechaje/{categoria_id}', 303)
